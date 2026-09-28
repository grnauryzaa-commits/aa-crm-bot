# -*- coding: utf-8 -*-
from docx import Document

doc = Document("reflections_kz.docx")

print("Всего абзацев:", len(doc.paragraphs))
print()
print("Первые 30 абзацев (repr — с видимыми спецсимволами):")
print("-" * 60)

for i, p in enumerate(doc.paragraphs[:30]):
    print(f"[{i:03}] {repr(p.text[:200])}")

print()
print("-" * 60)
print("Абзацы 100-130:")
print("-" * 60)
for i, p in enumerate(doc.paragraphs[100:130], start=100):
    print(f"[{i:03}] {repr(p.text[:200])}")